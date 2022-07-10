#!/usr/bin/env python3
# -*- encoding: utf-8 -*-
"""
__file__   : batch_rename
__time__   : 2022/5/15 00:53
__author__ : secoder
__resume__ : None
"""
import argparse
import os
import inspect
import shutil
import sys

from docx import Document


# def get_files(document_path):
#     files = os.listdir(document_path)
#     print(files)


class BatchRename(object):

    def __init__(self, path_name):
        self.path_name = path_name
        self.files = os.listdir(path_name)
        # return self.files

    def batch_rename(self):

        doc_files = self.files

        for doc_file in doc_files:
            print("读取的文件名：", doc_file)
            # 获取文件的当前路径
            path = os.path.join(self.path_name, doc_file)
            print("path" + path)
            document = Document(path)
            tables = document.tables
            # 从word中的第一个表读
            table = tables[0]
            table_num = table.cell(0, 1).text + '.docx'
            os.rename(path, table_num)
            print("重命名文件完成")

            new_files = os.listdir('./')
            for new_file in new_files:
                if new_file.endswith('.docx'):
                    shutil.move(new_file, self.path_name)


if __name__ == '__main__':
    docs_path_name = sys.argv[1]
    # docs_path_name = r'./docs'
    br = BatchRename(docs_path_name)
    br.batch_rename()
    # br.batch_rename()
    # print()
